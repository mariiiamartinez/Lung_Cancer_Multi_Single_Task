"""Group the raw cancer images by a label column of class_labels.xlsx
and build dataset/ (images + masks + labels.csv).

Every radiography in raw_dataset/cancer/ is linked to a patient via the ID in
its filename (01500603_PA_2V.jpg -> patient 1500603).
The patient label is read from the chosen grouping column of
raw_dataset/class_labels.xlsx and written as class_label in labels.csv. The
matching mask is copied from raw_dataset/masks_cancer/ by base name; images
whose patient has no label in the column, or that have no mask, are skipped.

Cancer and healthy images can also be grouped directly without using the
Excel file:
    python group_dataset.py cancer healthy

Healthy images can also be combined with cancer images grouped by an Excel
column:
    python group_dataset.py healthy "TYPE STAGE"

Usage:
    python group_dataset.py "TYPE STAGE"
    python group_dataset.py "CHT — NO_CHT"
    python group_dataset.py "UNIMODAL — MULTIMODAL"
    python group_dataset.py "INIT_CHT — INIT_SURGERY"
    python group_dataset.py cancer healthy
    python group_dataset.py healthy "TYPE STAGE"

Output:
    dataset/
    ├── images/         copies of the grouped radiographs
    ├── masks/          copies of their matching masks
    └── labels.csv      image_name + class_label per row
"""

import csv
import shutil
import sys
from pathlib import Path
from openpyxl import load_workbook


ROOT = Path('raw_dataset')
REGISTER = ROOT / 'class_labels.xlsx'
IMAGES_DIR = ROOT / 'cancer'
MASKS_DIR = ROOT / 'masks_cancer'
OUTPUT = Path('dataset')

IMAGE_EXTENSIONS = {'.jpg', '.jpeg'}
MASK_EXTENSION = '.png'


# Get command-line arguments
arguments = sys.argv[1:]

if not arguments:
    print('ERROR: pass a grouping column, e.g. python group_dataset.py "TYPE STAGE"')
    sys.exit(1)

# Dataset cleaning
if OUTPUT.exists():
    shutil.rmtree(OUTPUT)


# Find the specific column in Excel
def find_column(ws, name):
    for col, cell in enumerate(ws[1], start=1):
        if cell.value == name:
            return col
    return None


matched = []
no_mask = []
no_label = []
unmatched_masks = []


# Add images and masks directly without using Excel
def add_direct_group(group):
    images_dir = ROOT / group
    masks_dir = ROOT / f'masks_{group}'

    mask_stems = {}

    for mask in masks_dir.iterdir():
        if mask.is_file() and mask.suffix.lower() == MASK_EXTENSION:
            mask_stems.setdefault(mask.stem, mask.name)

    for img in sorted(images_dir.iterdir()):
        if not img.is_file() or img.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        mask_name = mask_stems.get(img.stem)

        if mask_name is None:
            no_mask.append(img.name)
            continue

        mask = masks_dir / mask_name
        matched.append((img, mask, group))

    all_image_stems = {
        img.stem for img in images_dir.iterdir()
        if img.is_file() and img.suffix.lower() in IMAGE_EXTENSIONS
    }

    unmatched_masks.extend([
        mask.name for mask in masks_dir.iterdir()
        if (
            mask.is_file()
            and mask.suffix.lower() == MASK_EXTENSION
            and mask.stem not in all_image_stems
        )
    ])


# Separate direct groups from Excel groupings
direct_groups = [
    arg.lower() for arg in arguments
    if arg.lower() in {'cancer', 'healthy'}
]

excel_groupings = [
    arg for arg in arguments
    if arg.lower() not in {'cancer', 'healthy'}
]

# Check arguments
if 'cancer' in direct_groups and excel_groupings:
    print('ERROR: cancer cannot be combined with an Excel grouping')
    sys.exit(1)


# Add direct cancer/healthy groups
for group in direct_groups:
    add_direct_group(group)


# Add cancer images grouped by Excel
if excel_groupings:
    if len(excel_groupings) != 1:
        print('ERROR: only one Excel grouping column can be used')
        sys.exit(1)

    grouping = excel_groupings[0]

    wb = load_workbook(REGISTER, read_only=True)
    ws = wb.active

    pid_col = find_column(ws, 'PATIENT ID')
    label_col = find_column(ws, grouping)

    if pid_col is None:
        print('ERROR: no PATIENT ID column in register')
        sys.exit(1)

    labels = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        pid = row[pid_col - 1]
        label = row[label_col - 1]

        if pid is not None and label is not None:
            labels[int(pid)] = str(label).strip()

    wb.close()

    mask_stems = {}

    for mask in MASKS_DIR.iterdir():
        if mask.is_file() and mask.suffix.lower() == MASK_EXTENSION:
            mask_stems.setdefault(mask.stem, mask.name)

    for img in sorted(IMAGES_DIR.iterdir()):
        if not img.is_file() or img.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        # Image filenames keep the leading zero, but Excel stores patient IDs as integers.
        pid = int(img.stem.split('_')[0])
        label = labels.get(pid)

        if label is None:
            no_label.append(img.name)
            continue

        mask_name = mask_stems.get(img.stem)

        if mask_name is None:
            no_mask.append(img.name)
            continue

        mask = MASKS_DIR / mask_name
        matched.append((img, mask, label))

    all_image_stems = {
        img.stem for img in IMAGES_DIR.iterdir()
        if img.is_file() and img.suffix.lower() in IMAGE_EXTENSIONS
    }

    unmatched_masks.extend([
        mask.name for mask in MASKS_DIR.iterdir()
        if (
            mask.is_file()
            and mask.suffix.lower() == MASK_EXTENSION
            and mask.stem not in all_image_stems
        )
    ])


# Set grouping description
if excel_groupings and direct_groups:
    grouping = f'healthy + {excel_groupings[0]}'
elif excel_groupings:
    grouping = excel_groupings[0]
else:
    grouping = ' + '.join(direct_groups)


images_out = OUTPUT / 'images'
masks_out = OUTPUT / 'masks'
images_out.mkdir(parents=True, exist_ok=True)
masks_out.mkdir(parents=True, exist_ok=True)


# Copy images and masks
for img, mask, label in matched:
    shutil.copy2(img, images_out / img.name)
    shutil.copy2(mask, masks_out / mask.name)


# Create labels CSV
with (OUTPUT / 'labels.csv').open('w', newline='') as fh:
    writer = csv.writer(fh)
    writer.writerow(['image_name', 'class_label'])

    for img, mask, label in matched:
        writer.writerow([img.name, label])


# Count images per class
counts = {}

for img, mask, label in matched:
    counts[label] = counts.get(label, 0) + 1


print(f'Grouping: {grouping}')
print(f'Created {OUTPUT / "labels.csv"}')
print(f'  Total images: {len(matched)}')
print(f'  Skipped: {len(no_mask)} no-mask + {len(no_label)} unlabeled')

if no_mask:
    print(f'    Images without mask: {no_mask}')

if no_label:
    print(f'    Images without label: {no_label}')

if unmatched_masks:
    print(f'  Masks without image: {unmatched_masks}')

print('  Per class:')

for label, count in sorted(counts.items(), key=lambda x: -x[1]):
    print(f'    {label}: {count}')

print(f'Images written to {images_out}')
print(f'Masks written to {masks_out}')
